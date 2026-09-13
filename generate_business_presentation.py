import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_business_pptx():
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Theme Color Palette
    COLOR_BG = RGBColor(11, 15, 25)          # Deep slate navy
    COLOR_CARD = RGBColor(18, 24, 38)        # Card background
    COLOR_BORDER = RGBColor(99, 102, 241)    # Indigo accent border
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252)# Pure white text
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Slate muted text
    COLOR_CYAN = RGBColor(6, 182, 212)       # Cyber cyan
    COLOR_VIOLET = RGBColor(139, 92, 246)    # Electric violet
    COLOR_EMERALD = RGBColor(16, 185, 129)   # Emerald green
    COLOR_ROSE = RGBColor(244, 63, 94)       # Rose red

    def apply_dark_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, subtitle_text, slide_num_str):
        # Header Box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10.5), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = 'Arial'
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_MAIN

        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.name = 'Arial'
        p2.font.size = Pt(14)
        p2.font.color.rgb = COLOR_TEXT_MUTED

        # Slide Number Pill
        num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.5), Inches(0.5), Inches(1.0), Inches(0.4))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = COLOR_CARD
        num_box.line.color.rgb = COLOR_CYAN
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = slide_num_str
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_CYAN

    # ==================== SLIDE 1: Title & Vision ====================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide1)

    # Main Hero Card
    hero_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.1))
    hero_card.fill.solid()
    hero_card.fill.fore_color.rgb = COLOR_CARD
    hero_card.line.color.rgb = COLOR_BORDER
    hero_card.line.width = Pt(1.5)

    tf1 = hero_card.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "\nOntoForge"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p_sub = tf1.add_paragraph()
    p_sub.text = "Enterprise Knowledge Graph & Automated Ontology Generation Engine\n"
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_TEXT_MAIN

    p_desc = tf1.add_paragraph()
    p_desc.text = (
        "Transforming raw enterprise databases and unstructured text into W3C-compliant OWL/RDF Ontologies "
        "and live Neo4j Knowledge Graphs with Gemini AI Integration."
    )
    p_desc.alignment = PP_ALIGN.CENTER
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = COLOR_TEXT_MUTED

    # ==================== SLIDE 2: Enterprise Challenge ====================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide2)
    add_header(slide2, "The Enterprise Data Challenge", "Siloed Schemas & Manual Ontology Construction Cost Organizations Millions", "02/08")

    challenges = [
        ("📁 80%+ Siloed Data", "Relational databases, CSVs, and PDFs lack explicit semantic topology, blocking graph analytics.", COLOR_ROSE),
        ("⏳ 6-Month Manual Effort", "Traditional W3C ontology modeling requires dedicated domain engineers and expensive manual labor.", COLOR_VIOLET),
        ("⚠️ RAG Hallucinations", "Enterprise LLMs produce hallucinations and broken queries without deterministic graph grounding.", COLOR_CYAN)
    ]

    for idx, (ch_title, ch_desc, ch_color) in enumerate(challenges):
        left_pos = Inches(0.8 + idx * 3.9)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.2), Inches(3.6), Inches(4.2))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = ch_color

        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = ch_title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ch_color

        p_body = tf.add_paragraph()
        p_body.text = f"\n{ch_desc}"
        p_body.font.size = Pt(13)
        p_body.font.color.rgb = COLOR_TEXT_MUTED

    # ==================== SLIDE 3: Solution Architecture ====================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide3)
    add_header(slide3, "End-to-End Solution Architecture", "Automated 4-Step Pipeline from Data Ingestion to Neo4j Graph Intelligence", "03/08")

    steps = [
        ("1. Data Ingestion", "SQL DBs, CSVs, Raw Text", COLOR_CYAN),
        ("2. Auto-Extraction", "OWL/RDF Class Matrix", COLOR_BORDER),
        ("3. Refine & Govern", "Cascading Class Deletion", COLOR_VIOLET),
        ("4. Graph & AI Sync", "Neo4j Cypher & Gemini", COLOR_EMERALD)
    ]

    for idx, (st_title, st_desc, st_color) in enumerate(steps):
        left_pos = Inches(0.8 + idx * 2.9)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.5), Inches(2.7), Inches(3.5))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = st_color

        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = st_title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = st_color

        p_b = tf.add_paragraph()
        p_b.text = f"\n{st_desc}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_TEXT_MUTED

    # ==================== SLIDE 4: App Screen - Class Matrix ====================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide4)
    add_header(slide4, "Application Showcase: Class Matrix Dashboard", "Interactive Workspace for Viewing and Editing Ontology Classes & Attributes", "04/08")

    app_window = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.8))
    app_window.fill.solid()
    app_window.fill.fore_color.rgb = COLOR_CARD
    app_window.line.color.rgb = COLOR_BORDER

    tf_app = app_window.text_frame
    tf_app.word_wrap = True

    p = tf_app.paragraphs[0]
    p.text = "🖥️ ONTOFORGE CLASS MATRIX DASHBOARD (Angular 19 & FastAPI UI)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p_content = tf_app.add_paragraph()
    p_content.text = (
        "\n• Visual Class Cards: Displays active ontology classes (e.g. Device, TemperatureSensor, Location).\n"
        "• Attribute Inspector: Real-time management of properties (id: string, temperature: float, status: string).\n"
        "• Direct Action Controls: 1-Click '🗑️ Delete Class' button with instant backend validation.\n"
        "• Graph Status Telemetry: Real-time connectivity indicator with target Neo4j instance."
    )
    p_content.font.size = Pt(14)
    p_content.font.color.rgb = COLOR_TEXT_MAIN

    # ==================== SLIDE 5: App Screen - Governance & Rebinding ====================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide5)
    add_header(slide5, "Application Showcase: Smart Schema Governance", "Cascading Class Deletion with Automatic Subclass Rebinding to owl:Thing", "05/08")

    gov_window = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.8))
    gov_window.fill.solid()
    gov_window.fill.fore_color.rgb = COLOR_CARD
    gov_window.line.color.rgb = COLOR_VIOLET

    tf_gov = gov_window.text_frame
    tf_gov.word_wrap = True

    p = tf_gov.paragraphs[0]
    p.text = "🛡️ CASCADING DEPENDENCY PROTECTION & REBINDING"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_VIOLET

    p_gov_b = tf_gov.add_paragraph()
    p_gov_b.text = (
        "\n1. Class Deletion Trigger: User clicks '🗑️ Delete Class' on target class 'Device'.\n"
        "2. Subclass Rebinding: Subclasses (e.g. TemperatureSensor) are automatically rebound to owl:Thing.\n"
        "3. Foreign Key Cleanup: Associated TargetGraphNode foreign keys & OntologyAttributes are safely purged.\n"
        "4. Zero Downstream Failure: Prevents graph corruption or broken query dependencies across applications."
    )
    p_gov_b.font.size = Pt(14)
    p_gov_b.font.color.rgb = COLOR_TEXT_MAIN

    # ==================== SLIDE 6: App Screen - Gemini Cypher Engine ====================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide6)
    add_header(slide6, "Application Showcase: Gemini AI GraphRAG", "Converting Natural Language Prompts into Deterministic Cypher Queries", "06/08")

    ai_window = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.8))
    ai_window.fill.solid()
    ai_window.fill.fore_color.rgb = COLOR_CARD
    ai_window.line.color.rgb = COLOR_EMERALD

    tf_ai = ai_window.text_frame
    tf_ai.word_wrap = True

    p = tf_ai.paragraphs[0]
    p.text = "🤖 GEMINI AI GRAPH INSIGHTS ENGINE"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_EMERALD

    p_ai_b = tf_ai.add_paragraph()
    p_ai_b.text = (
        "\n💬 User Prompt: 'Find all active temperature sensors reporting readings above 85 degrees.'\n\n"
        "⚡ Generated Cypher Query:\n"
        "   MATCH (s:TemperatureSensor)-[:LOCATED_IN]->(l:Location)\n"
        "   WHERE s.temperature > 85.0\n"
        "   RETURN s.id AS SensorID, s.temperature AS TempValue, l.facilityName AS Facility\n\n"
        "✔ Live Graph Response: Grounded in Neo4j live database records with zero hallucinations."
    )
    p_ai_b.font.size = Pt(14)
    p_ai_b.font.color.rgb = COLOR_TEXT_MAIN

    # ==================== SLIDE 7: Business Impact & ROI ====================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide7)
    add_header(slide7, "Enterprise Business Impact & ROI", "Quantifiable Time Savings, W3C Standards Compliance & Zero Schema Breakage", "07/08")

    metrics = [
        ("10x", "Faster Ontology Modeling", COLOR_CYAN),
        ("100%", "W3C OWL/RDF Compliant", COLOR_EMERALD),
        ("0%", "Broken Query Risks", COLOR_VIOLET),
        ("100%", "Deterministic GraphRAG", COLOR_BORDER)
    ]

    for idx, (m_val, m_lbl, m_color) in enumerate(metrics):
        left_pos = Inches(0.8 + idx * 2.9)
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.4), Inches(2.7), Inches(3.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = m_color

        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"\n{m_val}"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = m_color

        p_lbl = tf.add_paragraph()
        p_lbl.text = f"\n{m_lbl}"
        p_lbl.alignment = PP_ALIGN.CENTER
        p_lbl.font.size = Pt(13)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = COLOR_TEXT_MAIN

    # ==================== SLIDE 8: Deployment & Next Steps ====================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_background(slide8)
    add_header(slide8, "Enterprise Deployment & Next Steps", "Flexible Architecture Ready for Private Cloud and Microservices Integration", "08/08")

    dep_card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.5))
    dep_card.fill.solid()
    dep_card.fill.fore_color.rgb = COLOR_CARD
    dep_card.line.color.rgb = COLOR_CYAN

    tf_dep = dep_card.text_frame
    tf_dep.word_wrap = True

    p = tf_dep.paragraphs[0]
    p.text = "\n🚀 FLEXIBLE DEPLOYMENT OPTIONS"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p_options = tf_dep.add_paragraph()
    p_options.text = (
        "\n• On-Premises & Private Cloud: Single-command Docker container deployment.\n"
        "• Microservices Architecture: FastAPI REST API layer ready for enterprise API Gateway integration.\n"
        "• Dual Frontend Interface: Angular 19 Reactive Dashboard + Lightweight FastAPI Static Web App.\n"
        "• GitHub Repository: https://github.com/josephtijo1989/OntologyGenerator"
    )
    p_options.alignment = PP_ALIGN.LEFT
    p_options.font.size = Pt(14)
    p_options.font.color.rgb = COLOR_TEXT_MAIN

    out_path = r"c:\Users\TIJO\Documents\antigravity\quick-pasteur\OntoForge_Executive_Business_Pitch.pptx"
    prs.save(out_path)
    print(f"Executive presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    create_business_pptx()
